from django.contrib import admin
from simple_history.admin import SimpleHistoryAdmin
from django.contrib.auth import get_user_model
from django.http import HttpResponse
from django.urls import path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from datetime import datetime
from .models import HeroSlide, Statistics, Achievement, Ministry, MinistrySection, Gallery, Activity, DemographicData, SiteLogo, Subscriber, AboutPage, AboutMinistry, MissionVision, Challenge, BoardMember
from .models import Project, ProjectPage, VolunteerPage, GoTeam, GiveSection, PrayerPartner, VolunteerApplication
from .models import BlogPost, BlogImage, YouTubeVideo, SpotifyPodcast, MediaPage, DonationPage, BankAccount
from .models import Page, PageSection, PageSectionCard
from .models import FooterSettings, FooterLink, FooterSocial
from import_export.admin import ImportExportModelAdmin
from import_export.formats import base_formats
import xlsxwriter
from .resources import DemographicDataResource
from django.contrib.admin.models import LogEntry, ADDITION, CHANGE, DELETION
from django.template.response import TemplateResponse
from django.utils.html import format_html


##############AUDIT##################

class AuditAdminMixin(SimpleHistoryAdmin):
    def save_model(self, request, obj, form, change):
        if not change:
            obj.created_by = request.user
        obj.last_modified_by = request.user
        super().save_model(request, obj, form, change)

    readonly_fields = ('created_by', 'last_modified_by', 'created_at', 'updated_at')
    list_display = ('__str__', 'created_by', 'last_modified_by', 'updated_at')
##############AUDIT##################

class AuditAdmin(admin.ModelAdmin):
    list_display = ('action_time', 'user', 'content_type', 'object_repr', 'action_flag', 'change_message')
    list_filter = ('action_time', 'user', 'content_type', 'action_flag')
    search_fields = ('object_repr', 'change_message')
    date_hierarchy = 'action_time'
    ordering = ('-action_time',)
    readonly_fields = ('action_time', 'user', 'content_type', 'object_id', 'object_repr', 'action_flag', 'change_message')

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def action_flag_display(self, obj):
        flags = {
            ADDITION: 'Added',
            CHANGE: 'Changed',
            DELETION: 'Deleted',
        }
        return flags.get(obj.action_flag, '')
    action_flag_display.short_description = 'Action'

admin.site.register(LogEntry, AuditAdmin)


admin.site.register(SiteLogo)

@admin.register(HeroSlide)
class HeroSlideAdmin(AuditAdminMixin):
    list_display = ('title', 'order', 'is_active', 'created_at')
    list_editable = ('order', 'is_active')
    search_fields = ('title', 'description')
    list_filter = ('is_active',)
    ordering = ('order',)

@admin.register(Statistics)
class StatisticsAdmin(AuditAdminMixin):
    list_display = ('title', 'total_population', 'people_groups','states', 'lgas', 'wards', 'villages', 'converts', 'film_attendees', 'people_reached', 'updated_at')

@admin.register(Achievement)
class AchievementAdmin(AuditAdminMixin):
    list_display = ('villages', 'people_saved', 'people_reached', 'states_reached', 'lgas_reached', 'wards_reached', 'people_discipled', 'people_baptized', 'started_discipleship', 'bibles_shared', 'audiobibles_shared', 'sdcards_shared', 'updated_at')

    def has_add_permission(self, request):
        # Only allow one instance
        if self.model.objects.exists():
            return False
        return True


@admin.register(Ministry)
class MinistryAdmin(AuditAdminMixin):
    list_display = ('title', 'order')
    list_editable = ('order',)
    search_fields = ('title', 'description')
    ordering = ('order',)

@admin.register(Gallery)
class GalleryAdmin(AuditAdminMixin):
    list_display = ('title', 'is_featured', 'created_at')
    list_editable = ('is_featured',)
    search_fields = ('title', 'description')
    list_filter = ('is_featured',)


@admin.register(Activity)
class ActivityAdmin(AuditAdminMixin):
    list_display = ('title', 'order')
    list_editable = ('order',)
    search_fields = ('title',)
    ordering = ('order',)
    

@admin.register(MinistrySection)
class MinistrySectionAdmin(AuditAdminMixin):
    list_display = ('title', 'description', 'image')

    def has_add_permission(self, request):
        # Only allow one instance
        if self.model.objects.exists():
            return False
        return True


@admin.register(DemographicData)
class DemographicDataAdmin(ImportExportModelAdmin):
    resources_class = DemographicDataResource
    formats = [base_formats.XLSX]

    list_display = ['state', 'lga', 'ward', 'village', 'total_village_population', 'converts']
    list_filter = ['state', 'lga']
    search_fields = ['state', 'lga', 'ward', 'village']
    actions = ['export_selected_state']
    
    fieldsets = (
        ('Location', {
            'fields': ('state', 'lga', 'ward', 'village')
        }),
        ('Population Data', {
            'fields': (
                'christian_population',
                'muslim_population',
                'traditional_population',
                'converts',
                'total_village_population'
            )
        }),
        ('Additional Information', {
            'fields': ('film_attendance', 'people_group', 'practiced_religion')
        })
    )

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('download-template/', 
                 self.download_template, 
                 name='demographic-data-template'),
        ]
        return custom_urls + urls

    def download_template(self, request):
        # Create the HttpResponse object with Excel header
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="demographic_data_template.xlsx"'

        # Create the Excel workbook and add a worksheet
        workbook = xlsxwriter.Workbook(response)
        worksheet = workbook.add_worksheet()

        # Add headers
        headers = ['state', 'village', 'total_village_population', 'converts']
        for col, header in enumerate(headers):
            worksheet.write(0, col, header)

        # Add some example data (optional)
        example_data = [
            ['Lagos', 'Village A', 1000, 50],
            ['Abuja', 'Village B', 2000, 100],
        ]
        for row, data in enumerate(example_data, start=1):
            for col, value in enumerate(data):
                worksheet.write(row, col, value)

        workbook.close()
        return response

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context['show_template_download'] = True
        return super().changelist_view(request, extra_context)

    # def get_urls(self):
    #     urls = super().get_urls()
    #     custom_urls = [
    #         path('export-state/<str:state_name>/', 
    #              self.admin_site.admin_view(self.export_state_data), 
    #              name='export-state-data'),
    #     ]
    #     return custom_urls + urls

    def export_selected_state(self, request, queryset):
        # Get unique state from selection
        state_name = queryset.first().state if queryset.exists() else "Unknown"
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"{state_name} State Data"
        
        # Define headers
        headers = [
            'LGA', 'Ward', 'Village', 'Population', 'Christians',
            'Muslims', 'Traditional', 'Converts', 'Film Attendance'
        ]
        
        # Style headers
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        # Write data
        for row, data in enumerate(queryset, 2):
            ws.cell(row=row, column=1, value=data.lga)
            ws.cell(row=row, column=2, value=data.ward)
            ws.cell(row=row, column=3, value=data.village)
            ws.cell(row=row, column=4, value=data.total_village_population)
            ws.cell(row=row, column=5, value=data.christian_population)
            ws.cell(row=row, column=6, value=data.muslim_population)
            ws.cell(row=row, column=7, value=data.traditional_population)
            ws.cell(row=row, column=8, value=data.converts)
            ws.cell(row=row, column=9, value=data.film_attendance)
        
        # Adjust column widths
        for column in ws.columns:
            max_length = 0
            column = list(column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column[0].column_letter].width = adjusted_width
        
        # Create response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{state_name}_data_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        
        wb.save(response)
        return response
    
    export_selected_state.short_description = "Export selected data to Excel"


@admin.register(Subscriber)
class SubscriberAdmin(AuditAdminMixin):
    list_display = ('email', 'created_at', 'is_active')
    list_filter = ('is_active', 'created_at')
    search_fields = ('email',)
    date_hierarchy = 'created_at'

###### ABout ######
@admin.register(AboutPage)
class AboutPageAdmin(AuditAdminMixin):
    def has_add_permission(self, request):
        # Limit to single instance
        if self.model.objects.exists():
            return False
        return True

@admin.register(AboutMinistry)
class AboutMinistryAdmin(AuditAdminMixin):
    list_display = ('title', 'order')
    list_editable = ('order',)

@admin.register(MissionVision)
class MissionVisionAdmin(AuditAdminMixin):
    def has_add_permission(self, request):
        # Limit to single instance
        if self.model.objects.exists():
            return False
        return True

@admin.register(Challenge)
class ChallengeAdmin(AuditAdminMixin):
    list_display = ('title', 'order')
    list_editable = ('order',)
    fieldsets = (
        (None, {
            'fields': ('title', 'description', 'order')
        }),
        ('Challenge 1', {
            'fields': ('challenge1_title', 'challenge1_desc', 'challenge1_image')
        }),
        ('Challenge 2', {
            'fields': ('challenge2_title', 'challenge2_desc', 'challenge2_image')
        }),
        ('Challenge 3', {
            'fields': ('challenge3_title', 'challenge3_desc', 'challenge3_image')
        }),
    )

@admin.register(BoardMember)
class BoardMemberAdmin(AuditAdminMixin):
    list_display = ('name', 'position', 'order')
    list_editable = ('order',)

#########  PROJECTs ###########

@admin.register(ProjectPage)
class ProjectPageAdmin(AuditAdminMixin):
    def has_add_permission(self, request):
        if self.model.objects.exists():
            return False
        return True

@admin.register(Project)
class ProjectAdmin(AuditAdminMixin):
    list_display = ('title', 'order')
    list_editable = ('order',)

##############################

############  Volunteer ################
@admin.register(VolunteerPage)
class VolunteerPageAdmin(AuditAdminMixin):
    def has_add_permission(self, request):
        if self.model.objects.exists():
            return False
        return True

@admin.register(VolunteerApplication)
class VolunteerApplicationAdmin(AuditAdminMixin):
    list_display = ('full_name', 'email', 'area_of_interest', 'created_at')
    list_filter = ('area_of_interest', 'created_at')
    search_fields = ('full_name', 'email')
    readonly_fields = ('created_at',)


admin.site.register(GoTeam)
admin.site.register(PrayerPartner)
admin.site.register(GiveSection)

########################################

################# MEDIA ADMIN #########

class BlogImageInline(admin.TabularInline):
    model = BlogImage
    extra = 1

@admin.register(BlogPost)
class BlogPostAdmin(AuditAdminMixin):
    inlines = [BlogImageInline]
    list_display = ('title', 'author', 'published_date', 'is_featured', 'get_last_modified_by')
    list_filter = ('author', 'is_featured', 'published_date')
    readonly_fields = ('created_by', 'last_modified_by')
    
    def get_last_modified_by(self, obj):
        return obj.last_modified_by.username if obj.last_modified_by else '-'
    get_last_modified_by.short_description = 'Modified By'

    def save_model(self, request, obj, form, change):
        if not obj.pk:  # If this is a new object
            obj.created_by = request.user
        obj.last_modified_by = request.user
        super().save_model(request, obj, form, change)
        

@admin.register(YouTubeVideo)
class YouTubeVideoAdmin(AuditAdminMixin):
    list_display = ('title', 'published_date', 'is_featured')
    list_filter = ('is_featured', 'published_date')
    search_fields = ('title', 'description')

@admin.register(SpotifyPodcast)
class SpotifyPodcastAdmin(AuditAdminMixin):
    list_display = ('title', 'published_date', 'duration', 'is_featured')
    list_filter = ('is_featured', 'published_date')
    search_fields = ('title', 'description')

@admin.register(MediaPage)
class MediaPageAdmin(AuditAdminMixin):
    def has_add_permission(self, request):
        if self.model.objects.exists():
            return False
        return True
    
#########################################
@admin.register(DonationPage)
class DonationPageAdmin(AuditAdminMixin):
    list_display = ('title', 'bank_name')
    
    def has_add_permission(self, request):
        # Only allow one instance
        if self.model.objects.exists():
            return False
        return True

@admin.register(BankAccount)
class BankAccountAdmin(AuditAdminMixin):
    list_display = ('account_name', 'account_number', 'description')
    list_editable = ('account_number', 'description')
    search_fields = ('account_name', 'account_number')
    ordering = ('account_name',)


# ============================================================
#  WORDPRESS-LIKE PAGE BUILDER — Admin
# ============================================================

class PageSectionCardInline(admin.TabularInline):
    """Inline cards for a 'cards' section."""
    model  = PageSectionCard
    extra  = 2
    fields = ('order', 'icon', 'image', 'title', 'body', 'link_text', 'link_url')


class PageSectionInline(admin.StackedInline):
    """Inline content sections within a Page."""
    model   = PageSection
    extra   = 1
    ordering = ('order',)
    fields  = (
        ('section_type', 'order', 'is_visible'),
        ('background', 'css_class'),
        'title',
        'subtitle',
        'body',
        ('image', 'image_position'),
        ('button_text', 'button_url'),
        ('button2_text', 'button2_url'),
        'extra_data',
    )
    show_change_link = True


@admin.register(Page)
class PageAdmin(admin.ModelAdmin):
    list_display  = ('title', 'slug', 'show_in_nav', 'nav_order', 'is_published', 'updated_at')
    list_editable = ('show_in_nav', 'nav_order', 'is_published')
    prepopulated_fields = {'slug': ('title',)}
    search_fields = ('title', 'slug')
    list_filter   = ('is_published', 'show_in_nav')
    inlines       = [PageSectionInline]
    fieldsets = (
        ('Page Info', {
            'fields': ('title', 'slug', 'nav_label', 'meta_description', 'header_image')
        }),
        ('Navigation & Publishing', {
            'fields': (('is_published', 'show_in_nav', 'nav_order'),),
        }),
    )


@admin.register(PageSection)
class PageSectionAdmin(admin.ModelAdmin):
    list_display  = ('page', 'section_type', 'title', 'order', 'background', 'is_visible')
    list_editable = ('order', 'is_visible')
    list_filter   = ('page', 'section_type', 'background')
    search_fields = ('title', 'subtitle', 'body')
    inlines       = [PageSectionCardInline]
    fieldsets = (
        ('Location', {'fields': ('page', 'section_type', 'order', 'is_visible')}),
        ('Content', {'fields': ('title', 'subtitle', 'body', ('image', 'image_position'))}),
        ('Buttons', {'fields': (('button_text', 'button_url'), ('button2_text', 'button2_url'))}),
        ('Style', {'fields': ('background', 'css_class')}),
        ('Advanced (JSON)', {'fields': ('extra_data',), 'classes': ('collapse',)}),
    )

##############  SITE FOOTER  ##################

@admin.register(FooterSettings)
class FooterSettingsAdmin(AuditAdminMixin):
    list_display = ('__str__', 'email', 'phone', 'address')
    fieldsets = (
        ('Brand', {'fields': ('brand_text',)}),
        ('Contact', {'fields': ('contact_heading', 'email', ('phone', 'phone_alt'), 'address')}),
        ('Call To Action', {'fields': (('cta_text', 'cta_link'), 'show_cta')}),
        ('Column Headings', {'fields': (('quick_links_heading', 'involved_heading'),)}),
        ('Bottom Bar', {'fields': ('copyright_text', 'tagline')}),
    )

    def has_add_permission(self, request):
        # Singleton — only one footer configuration row.
        return not FooterSettings.objects.exists()


@admin.register(FooterLink)
class FooterLinkAdmin(admin.ModelAdmin):
    list_display  = ('label', 'column', 'url', 'order', 'is_active')
    list_editable = ('order', 'is_active')
    list_filter   = ('column', 'is_active')
    search_fields = ('label', 'url')


@admin.register(FooterSocial)
class FooterSocialAdmin(admin.ModelAdmin):
    list_display  = ('name', 'icon', 'url', 'order', 'is_active')
    list_editable = ('order', 'is_active')
    list_filter   = ('is_active',)
    search_fields = ('name', 'url')
